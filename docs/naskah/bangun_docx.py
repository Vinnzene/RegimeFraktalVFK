"""Bangun naskah skripsi .docx dari docs/naskah/isi_*.txt (urut nama berkas).

    python docs/naskah/bangun_docx.py <keluaran.docx>

Markup per baris (baris kosong memisahkan paragraf):
    #COVER                    halaman sampul
    #! JUDUL                  judul tanpa nomor (ABSTRAK, DAFTAR PUSTAKA)
    #TOC / #TOT / #TOF        daftar isi / tabel / gambar (field Word)
    #MAIN                     mulai bagian utama (nomor halaman 1, 2, 3, ...)
    # BAB 1|PENDAHULUAN       judul bab (halaman baru)
    ## 1.1. Judul             subbab;  ### 1.1.1. Judul  sub-subbab
    - butir / 1. butir        daftar
    TABEL: Tabel 1.1 ...      keterangan tabel, diikuti baris | a | b |
    SUMBER: ...               catatan sumber di bawah tabel/gambar
    GAMBAR: path | Gambar 3.1 ... | lebar_cm
    RUMUS: teks | (2.1)       persamaan bernomor
    PUSTAKA: entri            daftar pustaka (indentasi gantung)
    CATATAN: teks             catatan merah yang harus diisi/ditinjau penulis
    KATA: teks                baris kata kunci (tanpa indentasi)
Inline: **tebal**, *miring*.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SRC = sorted(Path(__file__).parent.glob("isi_*.txt"))
FONT = "Times New Roman"
LEBAR_TEKS = 14.0  # A4 21 cm − kiri 4 − kanan 3


def set_font(style_or_run, size=None, bold=None, italic=None, color=RGBColor(0, 0, 0)):
    f = style_or_run.font
    f.name = FONT
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(a), None)
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color


def siapkan_gaya(doc):
    st = doc.styles
    normal = st["Normal"]
    set_font(normal, 12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    for nama, size, align in (("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER),
                              ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT),
                              ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT)):
        s = st[nama]
        set_font(s, size, bold=True, italic=False)
        s.paragraph_format.alignment = align
        s.paragraph_format.space_before = Pt(12 if nama != "Heading 1" else 0)
        s.paragraph_format.space_after = Pt(6 if nama != "Heading 1" else 18)
        s.paragraph_format.line_spacing = 1.5
        s.paragraph_format.keep_with_next = True
    from docx.enum.style import WD_STYLE_TYPE
    for nama, align in (("Keterangan Tabel", WD_ALIGN_PARAGRAPH.CENTER),
                        ("Keterangan Gambar", WD_ALIGN_PARAGRAPH.CENTER)):
        s = st.add_style(nama, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        set_font(s, 12, bold=True)
        s.paragraph_format.alignment = align
        s.paragraph_format.line_spacing = 1.0
        s.paragraph_format.space_before = Pt(6)
        s.paragraph_format.space_after = Pt(6)
    s = st.add_style("Isi Tabel", WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = normal
    set_font(s, 10)
    s.paragraph_format.line_spacing = 1.0
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def halaman(sec):
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin, sec.top_margin = Cm(4), Cm(4)
    sec.right_margin, sec.bottom_margin = Cm(3), Cm(3)
    sec.footer_distance = Cm(1.5)


def nomor_halaman(sec, fmt):
    pg = sec._sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sec._sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), "1")
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, "PAGE", "1")


def field(p, instr, tampil=""):
    def fc(t):
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r = p.add_run()
        r._r.append(e)
    fc("begin")
    r = p.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fc("separate")
    p.add_run(tampil)
    fc("end")


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")


def tulis(p, teks, size=None, bold=None):
    for bagian in INLINE.split(teks):
        if not bagian:
            continue
        if bagian.startswith("**") and bagian.endswith("**"):
            r = p.add_run(bagian[2:-2])
            r.bold = True
        elif bagian.startswith("*") and bagian.endswith("*") and len(bagian) > 2:
            r = p.add_run(bagian[1:-1])
            r.italic = True
        else:
            r = p.add_run(bagian)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
    return p


def paragraf(doc, teks):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.0)
    return tulis(p, teks)


def butir(doc, penanda, teks, level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    kiri = 0.75 + 0.75 * level
    p.paragraph_format.left_indent = Cm(kiri)
    p.paragraph_format.first_line_indent = Cm(-0.75)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(kiri))
    p.add_run(f"{penanda}\t")
    return tulis(p, teks)


def shade(cell, warna):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), warna)
    tcpr.append(sh)


def tabel(doc, baris):
    rows = [[c.strip() for c in b.strip().strip("|").split("|")] for b in baris
            if not re.fullmatch(r"\|?\s*:?-{2,}.*", b.strip())]
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblw = t._tbl.tblPr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        t._tbl.tblPr.append(tblw)
    tblw.set(qn("w:type"), "pct")
    tblw.set(qn("w:w"), "5000")
    for i, r in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.paragraphs[0].style = doc.styles["Isi Tabel"]
            tulis(cell.paragraphs[0], r[j] if j < len(r) else "", bold=(i == 0) or None)
            if i == 0:
                shade(cell, "D9D9D9")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.append(h)
    sesudah = doc.add_paragraph()
    sesudah.paragraph_format.line_spacing = 1.0
    return t


def keterangan(doc, teks, gaya):
    p = doc.add_paragraph(style=gaya)
    m = re.match(r"((?:Tabel|Gambar) \d+\.\d+)\s+(.*)", teks)
    if m:
        p.add_run(m.group(1) + " ")
        r = p.add_run(m.group(2))
        r.bold = False
    else:
        p.add_run(teks)
    if gaya == "Keterangan Tabel":
        p.paragraph_format.keep_with_next = True
    return p


def sumber(doc, teks):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(8)
    tulis(p, teks, size=10)
    for r in p.runs:
        r.italic = True


def gambar(doc, path, ket, lebar):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ROOT / path), width=Cm(min(float(lebar), LEBAR_TEKS)))
    keterangan(doc, ket, "Keterangan Gambar")


def rumus(doc, teks, nomor):
    p = doc.add_paragraph()
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Cm(LEBAR_TEKS / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Cm(LEBAR_TEKS), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.add_run("\t")
    r = p.add_run(teks)
    r.font.name = "Cambria Math"
    r.font.size = Pt(11)
    p.add_run(f"\t{nomor}")


def catatan(doc, teks):
    p = doc.add_paragraph()
    r = p.add_run("[CATATAN PENULIS: " + teks + "]")
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def judul_bab(doc, teks):
    kepala, _, nama = teks.partition("|")
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.add_run(kepala.strip())
    if nama:
        p.add_run().add_break()
        p.add_run(nama.strip())


def judul_bebas(doc, teks, pertama=False):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = not pertama
    p.add_run(teks)


def daftar(doc, judul, instr, petunjuk):
    judul_bebas(doc, judul)
    p = doc.add_paragraph()
    field(p, instr, petunjuk)


def sampul(doc, meta):
    def baris(teks, size=12, bold=False, sebelum=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(sebelum)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(teks)
        r.bold = bold
        r.font.size = Pt(size)
    baris(meta["judul"], 14, True)
    baris("SKRIPSI", 14, True, 36)
    baris("Diajukan sebagai salah satu syarat untuk meraih", 12, False, 18)
    baris("Gelar Sarjana Komputer (S.Kom.)")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run("[Logo Universitas Bunda Mulia]")
    r.italic = True
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    baris("Oleh:", 12, False, 60)
    baris(meta["nama"], 12, True)
    baris(meta["nim"], 12, True)
    baris("UNIVERSITAS BUNDA MULIA", 12, True, 72)
    baris("FAKULTAS TEKNOLOGI DAN DESAIN", 12, True)
    baris("PROGRAM STUDI INFORMATIKA", 12, True)
    baris("JAKARTA", 12, True)
    baris("2026", 12, True)


def bangun(keluar: Path):
    doc = Document()
    siapkan_gaya(doc)
    halaman(doc.sections[0])
    meta = {"judul": "IDENTIFIKASI REGIME FRAKTAL DINAMIS BTC/USDT MULTI-TIMEFRAME MENGGUNAKAN "
                     "SWT DAN ATTENTION-BiLSTM BERBASIS HURST-DFA",
            "nama": "Vinnzene Fernando Karim", "nim": "32230115"}

    baris = [b for f in SRC for b in f.read_text(encoding="utf-8").splitlines() + [""]]
    i, buffer, pertama_depan = 0, [], True

    spasi = {"satu": False}

    def flush():
        if buffer:
            p = paragraf(doc, " ".join(s.strip() for s in buffer))
            if spasi["satu"]:
                p.paragraph_format.line_spacing = 1.0
            buffer.clear()

    while i < len(baris):
        s = baris[i].rstrip()
        i += 1
        if not s.strip() or s.startswith("%%"):
            flush()
            continue
        if s == "#COVER":
            flush()
            sampul(doc, meta)
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            halaman(sec)
            nomor_halaman(sec, "lowerRoman")
            continue
        if s == "#MAIN":
            flush()
            sec = doc.add_section(WD_SECTION.NEW_PAGE)
            halaman(sec)
            nomor_halaman(sec, "decimal")
            pertama_depan = True
            continue
        if s in ("#TOC", "#TOT", "#TOF"):
            flush()
            spasi["satu"] = False
            if s == "#TOC":
                daftar(doc, "DAFTAR ISI", 'TOC \\o "1-3" \\h \\z \\u', "Klik kanan → Update Field untuk menampilkan daftar isi.")
            elif s == "#TOT":
                daftar(doc, "DAFTAR TABEL", 'TOC \\h \\z \\t "Keterangan Tabel;1"', "Klik kanan → Update Field.")
            else:
                daftar(doc, "DAFTAR GAMBAR", 'TOC \\h \\z \\t "Keterangan Gambar;1"', "Klik kanan → Update Field.")
            pertama_depan = False
            continue
        if s.startswith("#! "):
            flush()
            judul_bebas(doc, s[3:].strip(), pertama=pertama_depan)
            spasi["satu"] = s[3:].strip() in ("ABSTRAK", "ABSTRACT")
            pertama_depan = False
            continue
        if s.startswith("### "):
            flush()
            doc.add_paragraph(s[4:].strip(), style="Heading 3")
            continue
        if s.startswith("## "):
            flush()
            doc.add_paragraph(s[3:].strip(), style="Heading 2")
            continue
        if s.startswith("# "):
            flush()
            if pertama_depan:
                kepala, _, nama = s[2:].partition("|")
                p = doc.add_paragraph(style="Heading 1")
                p.add_run(kepala.strip())
                p.add_run().add_break()
                p.add_run(nama.strip())
                pertama_depan = False
            else:
                judul_bab(doc, s[2:].strip())
            continue
        if s.startswith("TABEL:"):
            flush()
            keterangan(doc, s[6:].strip(), "Keterangan Tabel")
            rows = []
            while i < len(baris) and baris[i].strip().startswith("|"):
                rows.append(baris[i])
                i += 1
            tabel(doc, rows)
            continue
        if s.startswith("SUMBER:"):
            flush()
            sumber(doc, s[7:].strip())
            continue
        if s.startswith("GAMBAR:"):
            flush()
            path, ket, lebar = [x.strip() for x in s[7:].split("|")]
            gambar(doc, path, ket, lebar)
            continue
        if s.startswith("RUMUS:"):
            flush()
            teks, nomor = [x.strip() for x in s[6:].rsplit("|", 1)]
            rumus(doc, teks, nomor)
            continue
        if s.startswith("PUSTAKA:"):
            flush()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-1.0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(8)
            tulis(p, s[8:].strip())
            continue
        if s.startswith("CATATAN:"):
            flush()
            catatan(doc, s[8:].strip())
            continue
        if s.startswith("KATA:"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            tulis(p, s[5:].strip())
            continue
        m = re.match(r"^(\s*)(-|\d+\.|[a-z]\.)\s+(.*)", s)
        if m and not buffer:
            teks = m.group(3)
            while i < len(baris) and baris[i].startswith("   ") and not re.match(r"^\s*(-|\d+\.|[a-z]\.)\s", baris[i]):
                teks += " " + baris[i].strip()
                i += 1
            level = len(m.group(1)) // 3
            butir(doc, "•" if m.group(2) == "-" else m.group(2), teks, level)
            continue
        buffer.append(s)
    flush()

    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)
    doc.core_properties.author = meta["nama"]
    doc.core_properties.title = meta["judul"]
    doc.save(keluar)
    print("tersimpan:", keluar)


if __name__ == "__main__":
    bangun(Path(sys.argv[1]) if len(sys.argv) > 1 else Path("skripsi.docx"))
